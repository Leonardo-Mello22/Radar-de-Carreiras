"""Import the group's questionnaire and desk research without pooling populations."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd


FIELDS = {
    1: ("curso_graduacao", "Categoria de curso declarada."),
    2: ("periodo_graduacao_original", "Período declarado; desperiodizado não vira zero."),
    3: ("momento_carreira", "Momento de carreira autodeclarado, não senioridade inferida."),
    4: ("situacao_trabalho_declarada", "Resposta mistura regimes de trabalho e situação de busca; manter a categoria original."),
    5: ("facilidade_informacao_escala_1_5", "Nota ordinal. O XLSX não informa as âncoras da escala; não inferir sua direção."),
    6: ("dificuldade_principal", "Principal dificuldade declarada."),
    7: ("usaria_ferramenta_original", "Resposta original; itens fora de Sim/Não são sinalizados."),
    8: ("prioridades_painel_original", "Até três opções; vírgulas internas nas alternativas não são separadores."),
    9: ("expectativa_salarial_original", "Expectativa em reais mensais, não salário recebido."),
    10: ("busca_mudanca_original", "Busca atual de troca de emprego, área OU regime; não equivale ao alvo do State of Data."),
    11: ("cogitou_mudar_desistir_original", "Relato de ter cogitado mudar/desistir, não desistência efetivamente observada."),
    12: ("experiencia_ti_original", "Experiência declarada em meses ou texto; conversões mantêm os limites de intervalos."),
}

OPTIONS = {
    "salarios": "Faixas salariais detalhadas por cargo e nível (Júnior, Pleno, etc.)",
    "demanda_emprego": "Quais áreas e tecnologias possuem mais vagas ou admissões",
    "habilidades": "Requisitos técnicos e habilidades mais exigidos pelo mercado",
    "territorio": "Comparativo de mercado por região (ex: Pernambuco/Nordeste vs. outras regiões)",
    "vinculo": "Comparativo prático e financeiro entre tipos de vínculo (CLT vs. PJ/MEI)",
    "relatos_empresas": "Relatos ou avaliações de profissionais sobre as empresas",
    "transparencia": "Fontes oficiais de dados governamentais traduzidas de forma simples",
}


def clean(value):
    if pd.isna(value):
        return None
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value))).strip() or None


def integer(value):
    text = clean(value)
    if text is None or not re.fullmatch(r"\d+(?:\.0+)?", text):
        return None
    return int(float(text))


def yes_no(value):
    text = clean(value)
    return {"sim": True, "não": False, "nao": False}.get(text.casefold()) if text else None


def experience(value):
    """Return lower/upper months, keeping reported ranges non-exact."""
    text = clean(value)
    if text is None:
        return None, None, "ausente"
    number = integer(text)
    if number is not None:
        return number, number, "meses_numericos"
    text = text.casefold()
    text = re.sub(r"^um\b", "1", text)
    ranged = re.fullmatch(r"(\d+)\s+a\s+(\d+)\s+mes(?:es)?(?:\s+de\s+estágio)?", text)
    if ranged:
        low, high = map(int, ranged.groups())
        if low <= high:
            return low, high, "intervalo_declarado"
    duration = re.fullmatch(r"(?:(\d+)\s+anos?(?:\s+e\s+)?)?(?:(\d+)\s+mes(?:es)?)?", text)
    if duration and any(duration.groups()):
        years, months = (int(x or 0) for x in duration.groups())
        total = 12 * years + months
        return total, total, "texto_convertido"
    return None, None, "revisar_texto"


def priorities(value):
    text = clean(value)
    if not text:
        return [], "ausente"
    selected = []
    remainder = text
    for option, label in OPTIONS.items():
        if label in remainder:
            selected.append(option)
            remainder = remainder.replace(label, "", 1)
    if remainder.strip(" ,") or not 1 <= len(selected) <= 3:
        return [], "revisar_opcoes"
    return selected, "valido"


def export(frame, base):
    base.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(base.with_suffix(".csv"), index=False, encoding="utf-8-sig")
    frame.to_parquet(base.with_suffix(".parquet"), index=False)


def digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def prepare_questionnaire(path, output):
    frame = pd.read_excel(path, sheet_name=0, dtype=object)
    frame = frame.dropna(how="all")
    columns = {}
    for column in frame.columns:
        match = re.match(r"^\s*(\d+)\.", str(column))
        if match:
            qid = int(match.group(1))
            if qid in columns:
                raise ValueError(f"Repeated question number: {qid}")
            columns[qid] = column
    if set(columns) != set(FIELDS):
        raise ValueError("Questionnaire schema changed: review question mapping before importing.")
    times = pd.to_datetime(frame["Timestamp"], errors="raise")
    wide = pd.DataFrame({
        "respondente_id": [f"g12:{int(i) + 2:06d}" for i in frame.index],
        "linha_origem": [int(i) + 2 for i in frame.index],
        "fonte": "pesquisa_grupo12",
        "coleta_mes": times.dt.strftime("%Y-%m"),
    }, index=frame.index)
    mapping = []
    for qid, (name, description) in FIELDS.items():
        values = frame[columns[qid]].map(clean)
        wide[name] = pd.to_numeric(values, errors="coerce").astype("Int64") if qid == 5 else values.astype("string")
        if qid == 5 and not wide[name].dropna().between(1, 5).all():
            raise ValueError("Unexpected Likert scale value; do not silently alter its meaning.")
        mapping.append({"pergunta_id": f"Q{qid:02d}", "campo": name,
                        "pergunta_original": str(columns[qid]), "definicao": description})
    wide["periodo_graduacao_numero"] = wide.periodo_graduacao_original.str.extract(r"^(\d+)")[0].astype("Int64")
    for original, derived in [("usaria_ferramenta_original", "usaria_ferramenta"),
                              ("busca_mudanca_original", "busca_mudanca"),
                              ("cogitou_mudar_desistir_original", "cogitou_mudar_desistir")]:
        wide[derived] = wide[original].map(yes_no).astype("boolean")
        wide[f"{derived}_resposta_valida"] = wide[derived].notna()
    salary = wide.expectativa_salarial_original.map(integer).astype("Int64")
    wide["expectativa_salarial_mensal"] = salary
    wide["expectativa_salarial_revisar"] = salary.isna() | ~salary.between(300, 200000).fillna(False)
    exp = wide.experiencia_ti_original.map(experience)
    wide["experiencia_ti_meses_min"] = pd.array([x[0] for x in exp], dtype="Int64")
    wide["experiencia_ti_meses_max"] = pd.array([x[1] for x in exp], dtype="Int64")
    wide["experiencia_ti_meses_exatos"] = wide.experiencia_ti_meses_min.where(
        wide.experiencia_ti_meses_min.eq(wide.experiencia_ti_meses_max))
    wide["experiencia_ti_tratamento"] = pd.array([x[2] for x in exp], dtype="string")
    selection = wide.prioridades_painel_original.map(priorities)
    wide["prioridades_tratamento"] = pd.array([x[1] for x in selection], dtype="string")
    wide["prioridades_quantidade"] = pd.array([len(x[0]) if x[1] == "valido" else None for x in selection], dtype="Int64")
    long = wide.melt(id_vars=["respondente_id", "fonte"], value_vars=[v[0] for v in FIELDS.values()],
                     var_name="campo", value_name="resposta")
    long["resposta"] = long.resposta.map(clean).astype("string")
    long["pergunta_id"] = long.campo.map({v[0]: f"Q{k:02d}" for k, v in FIELDS.items()})
    items = pd.DataFrame([
        {"respondente_id": wide.loc[index, "respondente_id"], "pergunta_id": "Q08",
         "opcao_id": option, "opcao": OPTIONS[option]}
        for index, (options, _) in selection.items() for option in options
    ], columns=["respondente_id", "pergunta_id", "opcao_id", "opcao"])
    export(wide, output / "tratados/pesquisa_grupo12_respostas")
    export(long, output / "tratados/pesquisa_grupo12_respostas_longas")
    export(items, output / "tratados/pesquisa_grupo12_prioridades")
    export(pd.DataFrame(mapping), output / "auxiliares/questionario_grupo12")
    derived_definitions = {
        "respondente_id": "ID de linha estável, não comprova identidade de uma pessoa.",
        "linha_origem": "Linha do XLSX original, contando o cabeçalho como linha 1.",
        "fonte": "Identifica o questionário próprio, separado de outras populações.",
        "coleta_mes": "Mês da coleta; horário exato mantido somente no original privado.",
        "periodo_graduacao_numero": "Número inicial de período, sem imputar desperiodização.",
        "usaria_ferramenta": "Booleano apenas para Sim/Não; resposta fora do domínio vira ausente.",
        "busca_mudanca": "Busca de mudança de emprego, área ou regime declarada.",
        "cogitou_mudar_desistir": "Cogitou, não necessariamente realizou, mudança/desistência.",
        "expectativa_salarial_mensal": "Número inteiro declarado, preservado inclusive quando sinalizado.",
        "expectativa_salarial_revisar": "Ausente, não inteiro ou fora de 300 a 200000; revisão, não correção automática.",
        "experiencia_ti_meses_min": "Limite inferior de meses declarados; zero é uma resposta válida.",
        "experiencia_ti_meses_max": "Limite superior de meses declarados.",
        "experiencia_ti_meses_exatos": "Preenchido apenas quando limites coincidem; intervalos não viram média.",
        "experiencia_ti_tratamento": "Método de leitura: numérico, texto convertido, intervalo ou revisão.",
        "prioridades_tratamento": "Validação com as alternativas completas; não dividir em vírgulas internas.",
        "prioridades_quantidade": "Número de opções reconhecidas, no máximo três.",
    }
    definitions = {name: desc for name, desc in FIELDS.values()} | derived_definitions
    dictionary = pd.DataFrame([
        {"campo": c, "tipo": str(wide[c].dtype), "definicao": definitions.get(c, "Indicador de resposta pertencente ao domínio Sim/Não."),
         "ausentes": int(wide[c].isna().sum())} for c in wide.columns
    ])
    export(dictionary, output / "auxiliares/dicionario_pesquisa_grupo12")
    return {
        "registros": len(wide), "perguntas": len(FIELDS), "colunas_tratadas": len(wide.columns),
        "respostas_longas": len(long), "prioridades_selecionadas": len(items),
        "coleta_inicio": str(times.min().date()), "coleta_fim": str(times.max().date()),
        "q07_respostas_fora_dominio": int((wide.usaria_ferramenta_original.notna() & wide.usaria_ferramenta.isna()).sum()),
        "expectativas_sinalizadas": int(wide.expectativa_salarial_revisar.sum()),
        "experiencia_intervalos": int(wide.experiencia_ti_tratamento.eq("intervalo_declarado").sum()),
        "experiencia_textos_nao_convertidos": int(wide.experiencia_ti_tratamento.eq("revisar_texto").sum()),
        "prioridades_nao_reconhecidas": int(wide.prioridades_tratamento.eq("revisar_opcoes").sum()),
        "duplicatas_completas_originais": int(frame.duplicated().sum()),
        "deduplicacao": "Nenhuma linha removida; não há identificador de pessoa.",
        "privacidade": "Originais e respostas individuais locais, ignorados pelo Git; consentimento para divulgação não verificado.",
        "amostragem": "Método de recrutamento não documentado no XLSX; não assumir amostra probabilística ou representatividade.",
        "escala_q05": "Âncoras não disponíveis no XLSX; direção da escala precisa ser confirmada.",
        "alertas": ["Q07 tem respostas fora do domínio; não deslocar entre perguntas.",
                    "Expectativa sinalizada não é multiplicada por mil nem substituída.",
                    "Intervalo de experiência não vira um número exato.",
                    "Q10 não é intercambiável com planos_mudar_emprego_6m do State of Data."],
    }


def prepare_matrix(path, output):
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheet = book["Matriz"]
    fields = ["fonte_declarada", "topico", "achado_declarado", "confianca_declarada", "status_csd_declarado", "acao_proposta", "data_acesso_declarada"]
    records = []
    for row in sheet.iter_rows(min_row=2):
        values = [clean(c.value) for c in row[:7]]
        if not any(values):
            continue
        record = dict(zip(fields, values))
        record["referencia_id"] = f"desk:{row[0].row:04d}"
        record["url"] = row[0].hyperlink.target if row[0].hyperlink else None
        record["evidencia_verificada_nesta_integracao"] = False
        record["tipo_registro"] = "referencia_ou_hipotese"
        topic = (record["topico"] or "").casefold()
        warnings = {
            "progressão salarial": "Não medir progressão individual no CAGEDMOV: sem ligação entre pessoas, grupos de admissão e desligamento são diferentes.",
            "proxy de senioridade": "Não transformar idade/salário em senioridade observada; CAGEDMOV não fornece tempo de vínculo.",
            "rotatividade": "Fonte secundária; conferir população, período e levantamento original antes de usar o percentual.",
            "soluções similares": "A afirmação sobre ausência de concorrentes não foi verificada nesta integração.",
        }
        record["alerta_metodologico"] = next((text for key, text in warnings.items() if key in topic),
                                              "Conteúdo transcrito do grupo; importar não confirma a evidência externa.")
        records.append(record)
    result = pd.DataFrame(records)
    export(result, output / "tratados/matriz_pesquisa_grupo12")
    return {"registros": len(result), "links_preservados": int(result.url.notna().sum()),
            "status_csd_declarados": result.status_csd_declarado.value_counts().to_dict(),
            "verificacao": "Status CSD e confiança preservados como declaração do grupo, não como validação automática."}


def validate_group12(output):
    manifest = json.loads((output / "fontes_pesquisa_grupo12.json").read_text(encoding="utf-8"))
    for source in manifest["arquivos"]:
        assert digest(output / source["arquivo"]) == source["sha256"]
    wide = pd.read_parquet(output / "tratados/pesquisa_grupo12_respostas.parquet")
    long = pd.read_parquet(output / "tratados/pesquisa_grupo12_respostas_longas.parquet")
    items = pd.read_parquet(output / "tratados/pesquisa_grupo12_prioridades.parquet")
    original = pd.read_excel(output / "brutos/pesquisa_grupo12/trajetoria_profissional_respostas.xlsx", dtype=object).dropna(how="all")
    assert len(wide) == len(original) == manifest["questionario"]["registros"]
    assert wide.respondente_id.is_unique and wide.linha_origem.is_unique
    assert "Timestamp" not in wide.columns
    assert len(long) == len(wide) * len(FIELDS)
    assert not long.duplicated(["respondente_id", "pergunta_id"]).any()
    assert set(long.respondente_id) == set(wide.respondente_id)
    assert set(items.respondente_id) <= set(wide.respondente_id)
    assert not items.duplicated(["respondente_id", "opcao_id"]).any()
    assert set(items.opcao_id) <= set(OPTIONS)
    counts = items.groupby("respondente_id").size().reindex(wide.respondente_id, fill_value=0)
    assert counts.to_list() == wide.prioridades_quantidade.fillna(0).to_list()
    assert counts.between(0, 3).all()
    for header in original:
        match = re.match(r"^\s*(\d+)\.", str(header))
        if not match:
            continue
        qid = int(match.group(1))
        name = FIELDS[qid][0]
        expected = original[header].map(clean).tolist()
        actual = wide[name].map(clean).tolist()
        assert actual == expected, name
    for original_name, boolean in [("usaria_ferramenta_original", "usaria_ferramenta"),
                                   ("busca_mudanca_original", "busca_mudanca"),
                                   ("cogitou_mudar_desistir_original", "cogitou_mudar_desistir")]:
        expected = wide[original_name].map(yes_no).astype("boolean")
        pd.testing.assert_series_equal(wide[boolean], expected, check_names=False)
    exact = wide.experiencia_ti_meses_min.eq(wide.experiencia_ti_meses_max).fillna(False)
    assert wide.loc[~exact, "experiencia_ti_meses_exatos"].isna().all()
    assert wide.loc[exact, "experiencia_ti_meses_exatos"].eq(wide.loc[exact, "experiencia_ti_meses_min"]).all()
    assert int(wide.expectativa_salarial_revisar.sum()) == manifest["questionario"]["expectativas_sinalizadas"]
    matrix = pd.read_parquet(output / "tratados/matriz_pesquisa_grupo12.parquet")
    assert len(matrix) == manifest["matriz"]["registros"]
    assert matrix.referencia_id.is_unique
    assert not matrix.evidencia_verificada_nesta_integracao.any()
    sheet = openpyxl.load_workbook(output / "brutos/pesquisa_grupo12/matriz_pesquisa_g12.xlsx", data_only=True)["Matriz"]
    for record in matrix.itertuples():
        row = int(record.referencia_id.split(":")[-1])
        expected = sheet.cell(row, 1).hyperlink
        assert (record.url if pd.notna(record.url) else None) == (expected.target if expected else None)
        assert record.achado_declarado == clean(sheet.cell(row, 3).value)
    checks = ["SHA-256 dos dois XLSX conferido", f"{len(wide)} respostas conciliadas com as {len(FIELDS)} perguntas originais",
              "Tabela longa e seleção de prioridades reconciliadas sem dividir vírgulas internas",
              "Sim/Não, salário sinalizado e experiência sem imputação conferidos",
              "Matriz reconciliada com os achados e links originais", "Timestamp exato ausente da base analítica"]
    result = {"status": "APROVADO_ESTRUTURA_COM_ALERTAS", "verificacoes": checks,
              "questionario": manifest["questionario"], "matriz": manifest["matriz"]}
    (output / "qualidade/validacao_pesquisa_grupo12.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return checks


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--responses", type=Path, required=True)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    raw = args.output / "brutos/pesquisa_grupo12"
    raw.mkdir(parents=True, exist_ok=True)
    sources = []
    for source, name, kind in [(args.responses, "trajetoria_profissional_respostas.xlsx", "questionario_proprio"),
                               (args.matrix, "matriz_pesquisa_g12.xlsx", "desk_research")]:
        dest = raw / name
        if source.resolve() != dest.resolve():
            shutil.copy2(source, dest)
        if digest(source) != digest(dest):
            raise ValueError("Source copy checksum mismatch")
        sources.append({"arquivo": f"brutos/pesquisa_grupo12/{name}", "tipo": kind,
                        "sha256": digest(dest), "bytes": dest.stat().st_size,
                        "origem": "XLSX fornecido pelo Grupo 12; não obtido de fonte pública."})
    questionnaire = prepare_questionnaire(raw / "trajetoria_profissional_respostas.xlsx", args.output)
    matrix = prepare_matrix(raw / "matriz_pesquisa_g12.xlsx", args.output)
    manifest = {"arquivos": sources, "questionario": questionnaire, "matriz": matrix}
    (args.output / "fontes_pesquisa_grupo12.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output / "qualidade").mkdir(exist_ok=True)
    (args.output / "qualidade/relatorio_pesquisa_grupo12.json").write_text(
        json.dumps({"status": "IMPORTADO_COM_ALERTAS", "questionario": questionnaire, "matriz": matrix},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"questionario": questionnaire, "matriz": matrix}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
